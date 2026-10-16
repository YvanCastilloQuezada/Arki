from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

repo = Path(__file__).resolve().parents[1]
out = repo / 'docs' / 'arki-research' / 'ARKI_WEB_RESEARCH_STAGING.xlsx'
out.parent.mkdir(parents=True, exist_ok=True)

wb = Workbook()
readme = wb.active
readme.title = '00_README'
readme.append(['ARKI — Web Research Superpower — STAGING'])
readme.append(['Purpose', 'Receive abstract metadata and vocabulary candidates from public web research before human review.'])
readme.append(['Canonical rule', 'Nothing in this workbook is a FACT or canonical Core knowledge until human validation is recorded.'])
readme.append(['Rights rule', 'Store metadata and short abstracts only; do not reproduce copyrighted source text.'])
readme.append(['Automation rule', 'Daily cycle is limited to 30 sources and requires owner human authorization; candidates remain pending.'])
readme.append(['Status values', 'DISCOVERED | ABSTRACTED | STAGED_FOR_EXCEL | HUMAN_CONFIRMED | REJECTED | UNKNOWN'])

queue = wb.create_sheet('01_REVIEW_QUEUE')
queue.append(['SOURCE_ID', 'AUTHOR_OR_INSTITUTION', 'TITLE', 'URL', 'SOURCE_TYPE', 'TOPIC', 'ABSTRACT', 'TERMS', 'RIGHTS', 'STAGE', 'EPISTEMIC_STATUS', 'HUMAN_REVIEW_REQUIRED', 'CAPTURED_AT', 'CONTENT_HASH', 'NOTES'])

vocab = wb.create_sheet('02_VOCABULARY_CANDIDATES')
vocab.append(['TERM_ID', 'TERM', 'DEFINITION_CANDIDATE', 'SOURCE_ID', 'CONTEXT', 'LANGUAGE', 'STATUS', 'HUMAN_REVIEW_REQUIRED', 'NOTES'])

log = wb.create_sheet('03_CHANGE_LOG')
log.append(['CHANGE_ID', 'AT', 'ACTOR', 'ACTION', 'TARGET_SHEET', 'TARGET_ID', 'REASON', 'PREVIOUS_STATUS', 'NEW_STATUS', 'PROVENANCE'])

for ws in wb.worksheets:
    ws.freeze_panes = 'A2'
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='0B2945')
        cell.alignment = Alignment(wrap_text=True, vertical='top')
    for column in ws.columns:
        letter = column[0].column_letter
        ws.column_dimensions[letter].width = min(max(max(len(str(c.value or '')) for c in column) + 2, 14), 42)

wb.save(out)
print(out)
